import copy
import io
import re
from pathlib import Path
from xml.sax.saxutils import escape
from .dates import task_state


def anonymize(item):
    item = copy.deepcopy(item)
    names = list(dict.fromkeys([*item['speakers'].values(), *(t['assignee'] for t in item['tasks'] if t['assignee'])]))
    replacements = {name: f'Участник {i + 1}' for i, name in enumerate(names)}
    def scrub(value):
        if isinstance(value, str):
            for name in sorted(replacements, key=len, reverse=True):
                value = value.replace(name, replacements[name])
            return value
        if isinstance(value, list):
            return [scrub(v) for v in value]
        if isinstance(value, dict):
            return {k: scrub(v) for k, v in value.items()}
        return value
    item = scrub(item)
    item['title'] = 'Совещание — обезличенная копия'
    item['filename'] = ''
    item['anonymized'] = True
    return item


def timestamp(seconds):
    return f'{int(seconds) // 60:02d}:{int(seconds) % 60:02d}'


def rows(item):
    statuses = {'todo': 'К исполнению', 'in_progress': 'В работе', 'done': 'Выполнено', 'overdue': 'Просрочено'}
    for index, task in enumerate(item['tasks'], 1):
        due = task.get('due_date') or 'Не определён'
        if task.get('date_uncertain'):
            due += ' (уточнить)'
        yield [str(index), task['title'], task['assignee'] or 'Не назначен', due,
               statuses[task_state(task)] + ('; проверено' if task.get('reviewed') else '; черновик')]


def docx_bytes(item):
    from docx import Document
    from docx.shared import Pt, Cm
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(10)
    doc.add_heading('ПРОТОКОЛ СОВЕЩАНИЯ', 0)
    doc.add_heading(item['title'], 1)
    doc.add_paragraph(f"Дата: {item['date']} • Alem AI • Локальная обработка")
    doc.add_paragraph('Черновик, сформированный ИИ. Непроверенные поручения и неоднозначные сроки требуют подтверждения секретаря.')
    if item.get('anonymized'):
        doc.add_paragraph('Известные имена заменены. Перед публичной демонстрацией проверьте весь текст на косвенные идентификаторы.')
    doc.add_heading('Краткое содержание', 1)
    doc.add_paragraph(item.get('summary', ''))
    doc.add_heading('Решения', 1)
    for decision in item.get('decisions', []):
        doc.add_paragraph(decision, style='List Bullet')
    doc.add_heading('Поручения', 1)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Light Shading Accent 1'
    for cell, value in zip(table.rows[0].cells, ['№', 'Поручение', 'Ответственный', 'Срок', 'Статус']):
        cell.text = value
    for row in rows(item):
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value
    doc.add_heading('Основания поручений', 1)
    for i, task in enumerate(item['tasks'], 1):
        doc.add_paragraph(f"{i}. Реплики {', '.join(map(str, task['evidence_ids']))}. {task['evidence']}")
    doc.add_heading('Транскрипт', 1)
    for s in item['segments']:
        name = item['speakers'].get(s['speaker'], s['speaker'])
        p = doc.add_paragraph()
        p.add_run(f"[{s['id']}] {timestamp(s['start'])} · {name}: ").bold = True
        p.add_run(s['text'])
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def pdf_bytes(item):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    choices = [Path(__file__).parents[1] / 'assets' / 'DejaVuSans.ttf', Path('C:/Windows/Fonts/arial.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
    font = next((p for p in choices if p.exists()), None)
    if font is None:
        raise RuntimeError('Для PDF требуется DejaVuSans.ttf или Arial с кириллицей.')
    if 'Alem' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('Alem', str(font)))
    styles = getSampleStyleSheet()
    for key in ['Normal', 'Heading1', 'Heading2', 'Title']:
        styles[key].fontName = 'Alem'
    styles['Normal'].fontSize = 9
    styles['Normal'].leading = 14
    small = ParagraphStyle('small', parent=styles['Normal'], fontSize=8, leading=11)
    def p(text, style='Normal'):
        return Paragraph(escape(str(text)).replace('\n', '<br/>'), styles[style])
    story = [p('ПРОТОКОЛ СОВЕЩАНИЯ', 'Title'), p(item['title'], 'Heading1'), p(f"{item['date']} · Alem AI · Локальная обработка"), Spacer(1, 12),
             p('Черновик ИИ. Подтвердите поручения и неоднозначные сроки перед использованием.'), p('Краткое содержание', 'Heading2'), p(item.get('summary', ''))]
    if item.get('anonymized'):
        story.append(p('Известные имена заменены. Проверьте текст на косвенные идентификаторы перед публикацией.'))
    story.append(p('Решения', 'Heading2'))
    story.extend(p('• ' + d) for d in item.get('decisions', []))
    story.append(p('Поручения', 'Heading2'))
    table_data = [[Paragraph(escape(c), small) for c in r] for r in [['№', 'Поручение', 'Ответственный', 'Срок', 'Статус'], *rows(item)]]
    table = Table(table_data, colWidths=[22, 190, 100, 92, 95], repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e3eee9')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#d6ddd9')), ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6), ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
    story.append(table)
    story.append(p('Основания поручений', 'Heading2'))
    for index, task in enumerate(item['tasks'], 1):
        story.append(p(f"{index}. Реплики {', '.join(map(str, task['evidence_ids']))}: {task['evidence']}"))
    story.append(p('Транскрипт', 'Heading2'))
    for s in item['segments']:
        story.extend([p(f"[{s['id']}] {timestamp(s['start'])} · {item['speakers'].get(s['speaker'], s['speaker'])}"), p(s['text']), Spacer(1, 7)])
    def footer(canvas, doc):
        canvas.setFont('Alem', 8)
        canvas.drawString(48, 25, 'Alem AI · Протокол совещания')
        canvas.drawRightString(A4[0] - 48, 25, str(doc.page))
    stream = io.BytesIO()
    SimpleDocTemplate(stream, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=40, bottomMargin=44, title='Протокол совещания').build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
